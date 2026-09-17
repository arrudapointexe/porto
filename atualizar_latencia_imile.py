import os
import re
import sys
import unicodedata
import pandas as pd
from playwright.sync_api import sync_playwright
from dotenv import load_dotenv
from openpyxl import load_workbook

# Fix para o terminal do Windows não travar com emojis
sys.stdout.reconfigure(encoding='utf-8')

# Carrega as variáveis de ambiente do arquivo .env ANTES de qualquer outra coisa
load_dotenv()

# Meus módulos locais, baseados nos seus scripts existentes
from imile_utils import login_imile
from config import BASES_SLA, ITR_LOC_CITIES, ITR_INT_CITIES, JML_LOC_CITIES, JML_INT_CITIES, SNB_LOC_CITIES, SNB_INT_CITIES

# ==========================================
# CONFIGURAÇÕES
# ==========================================
# Mapeia a sigla da base para o MESMO arquivo de latência final consolidado
MAPA_ARQUIVOS_LATENCIA = {
    "JML": "latencia_consolidada.xlsx",
    "ITR": "latencia_consolidada.xlsx",
    "SNB": "latencia_consolidada.xlsx"
}

# ==========================================
# FUNÇÃO AUXILIAR PARA FORMATAÇÃO
# ==========================================
def preparar_colunas_para_arquivo_excel(df):
    """
    Converte colunas de código/waybill para texto antes de exportar para Excel.
    Isso evita notação científica e preserva valores longos como números de tracking.
    """
    if df is None:
        return None

    df = df.copy()
    colunas_alvo = []
    for coluna in df.columns:
        nome = str(coluna).strip().lower()
        if any(token in nome for token in ['waybill', 'tracking', 'order', 'pack', 'codigo', 'code', 'awb', 'bill']):
            colunas_alvo.append(coluna)

    for coluna in colunas_alvo:
        df[coluna] = df[coluna].apply(lambda valor: '' if pd.isna(valor) else str(valor).strip()).astype(object)

    return df


def normalizar_cidade(texto):
    if pd.isna(texto):
        return ""
    texto = unicodedata.normalize('NFKD', str(texto))
    texto = ''.join(char for char in texto if not unicodedata.combining(char))
    texto = texto.encode('ascii', 'ignore').decode('ascii')
    texto = texto.strip().lower()
    texto = re.sub(r'[^a-z0-9\s]', ' ', texto)
    texto = re.sub(r'\s+', ' ', texto).strip()
    return texto


def adicionar_coluna_abrangencia(df):
    """
    Adiciona a coluna 'abrangencia' com 1 para cidades da abrangência e 0 para as demais.
    Usa as listas do config.py como base para as cidades LOC e INT de cada base.
    """
    if df is None:
        return None

    df = df.copy()
    df['abrangencia'] = 0

    col_city = buscar_coluna_por_tokens(df.columns, ['consignee city', 'city'])
    col_base = buscar_coluna_por_tokens(df.columns, ['base'])

    if col_city is None or col_base is None:
        return df

    cidades_abrangencia = {
        'ITR': [*ITR_LOC_CITIES, *ITR_INT_CITIES],
        'JML': [*JML_LOC_CITIES, *JML_INT_CITIES],
        'SNB': [*SNB_LOC_CITIES, *SNB_INT_CITIES],
    }

    cidades_norm = {
        base: [normalizar_cidade(cidade) for cidade in cidades]
        for base, cidades in cidades_abrangencia.items()
    }

    def marcar_abrangencia(base_valor, cidade_valor):
        base = str(base_valor or '').strip().upper()
        cidade = normalizar_cidade(cidade_valor)
        if not cidade:
            return 0

        cidades_base = cidades_norm.get(base, [])
        if cidade in cidades_base:
            return 1

        for cidade_base in cidades_base:
            if cidade.startswith(cidade_base) or cidade_base.startswith(cidade):
                return 1

        return 0

    df['abrangencia'] = df.apply(lambda row: marcar_abrangencia(row[col_base], row[col_city]), axis=1)
    return df


def formatar_waybill_como_texto(caminho_arquivo):
    """
    Formata todas as colunas de waybill/tracking como texto no Excel para evitar notação científica.
    """
    try:
        wb = load_workbook(caminho_arquivo)
        ws = wb.active
        
        # Itera pelas colunas do header para identificar waybill/tracking
        for col_idx, col_cell in enumerate(ws[1], 1):
            if col_cell.value and any(x in str(col_cell.value).lower() for x in ['waybill', 'tracking', 'order', 'pack']):
                # Formata todas as células dessa coluna como texto
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                    for cell in row:
                        cell.number_format = '@'  # @ = texto no Excel
                        if cell.value is not None:
                            cell.value = str(cell.value)  # Converte para string
        
        wb.save(caminho_arquivo)
        print(f"INFO: Arquivo '{os.path.basename(caminho_arquivo)}' formatado com waybill como texto.")
    except Exception as e:
        print(f"WARN: Erro ao formatar arquivo: {e}")

# ==========================================
# FUNÇÕES DE ENRIQUECIMENTO POR DESTINATÁRIO
# ==========================================

def normalizar_texto(texto):
    if pd.isna(texto):
        return ""
    return str(texto).strip().lower()


def buscar_coluna_por_tokens(colunas, tokens):
    tokens = [t.lower() for t in tokens]
    for col in colunas:
        nome = normalizar_texto(col)
        for token in tokens:
            if token in nome:
                return col
    return None


def identificar_colunas_destinatario(df):
    mapping = {
        'Consignee Province': ['consignee province', 'province'],
        'Consignee City': ['consignee city', 'city'],
        'Consignee Area': ['consignee area', 'area', 'district', 'subdistrict', 'county'],
        'Consignee Address': ['consignee address', 'address'],
        'Uploaded Declared Value': ['uploaded declared value', 'declared value', 'uploaded declared', 'declared'],
        'Last Problem Type': ['last problem type', 'problem type'],
        'Last Problem Reason': ['last problem reason', 'problem reason']
    }

    result = {}
    for nome, candidatos in mapping.items():
        coluna = buscar_coluna_por_tokens(df.columns, candidatos)
        if coluna:
            result[nome] = coluna
    return result


def limpar_html_para_texto(texto):
    import html as html_module
    texto = html_module.unescape(texto or "")
    texto = re.sub(r'<script.*?>.*?</script>', '', texto, flags=re.S | re.I)
    texto = re.sub(r'<style.*?>.*?</style>', '', texto, flags=re.S | re.I)
    texto = re.sub(r'<[^>]+>', '', texto)
    return texto.strip()


def parse_html_table_simple(html_text):
    tabelas = []
    for table_html in re.findall(r'(<table[\s\S]*?>[\s\S]*?</table>)', html_text, flags=re.I):
        linhas = []
        for row_html in re.findall(r'<tr[\s\S]*?>([\s\S]*?)</tr>', table_html, flags=re.I):
            celulas = []
            for cell_html in re.findall(r'<t[dh][\s\S]*?>([\s\S]*?)</t[dh]>', row_html, flags=re.I):
                texto = limpar_html_para_texto(cell_html)
                celulas.append(texto)
            if celulas:
                linhas.append(celulas)
        if linhas:
            if len(linhas) > 1 and len(set(len(l) for l in linhas)) == 1:
                df = pd.DataFrame(linhas[1:], columns=linhas[0])
            else:
                df = pd.DataFrame(linhas)
            tabelas.append(df)
    return tabelas


def encontrar_tabela_destinatario(df_list, sigla):
    if not df_list:
        return None
    if len(df_list) == 1:
        return df_list[0]
    melhores = sorted(df_list, key=lambda df: (df.shape[1] * df.shape[0], df.shape[1]), reverse=True)
    for df in melhores:
        cabecalhos = ' '.join(map(str, df.columns)).lower()
        if any(token in cabecalhos for token in ['consignee', 'uploaded', 'declared', 'area', 'address']):
            return df
    return melhores[0]


def carregar_tabela_html_para_df(page, sigla):
    try:
        html = page.content()
        tabelas = parse_html_table_simple(html)
        if not tabelas:
            print(f"[{sigla}] WARN: Nenhuma tabela HTML encontrada no conteúdo da página.")
            return None

        return encontrar_tabela_destinatario(tabelas, sigla)
    except Exception as e:
        print(f"[{sigla}] WARN: Falha ao ler tabela do HTML: {e}")
        return None


def dividir_em_lotes(lista, tamanho_lote=1000):
    for indice in range(0, len(lista), tamanho_lote):
        yield lista[indice:indice + tamanho_lote]


def buscar_arquivo_destinatario(page, awb_list, sigla):
    if not awb_list:
        print(f"[{sigla}] WARN: Lista de AWBs vazia. Nenhum dado de destinatário será extraído.")
        return None

    url = "https://ds.imile.com/#/DSOperation/WaybillManagement/dispatchWaybillQueryDs"
    print(f"[{sigla}] INFO: Navegando para pesquisa por destinatário...")
    page.goto(url, wait_until="domcontentloaded")
    page.wait_for_timeout(5000)

    # Fallback de navegação via menu caso o hash direto não carregue a tela esperada
    try:
        if page.get_by_text("Dispatch waybill query").count() > 0:
            page.get_by_text("Dispatch waybill query").first.click(force=True)
            page.wait_for_timeout(4000)
    except Exception:
        pass

    def encontrar_input(page_to_search):
        selectors = [
            'input[placeholder*="Por favor"]',
            'input[placeholder*="por favor"]',
            'input[placeholder*="Please"]',
            'input[placeholder*="please"]',
            'input[placeholder*="insira"]',
            'input[placeholder*="insert"]',
            'div.MuiInputBase-root input[type="text"]',
            'input[type="search"]',
            'input[type="text"]'
        ]
        for selector in selectors:
            try:
                locator = page_to_search.locator(selector).first
                if locator.count() and locator.is_visible():
                    return locator
            except Exception:
                continue
        return None

    def consultar_lote(awb_lote, indice_lote):
        input_locator = encontrar_input(page)
        if input_locator is None:
            print(f"[{sigla}] INFO: Campo de AWB não encontrado no conteúdo principal. Tentando menu lateral...")
            try:
                page.get_by_text("Waybill management").first.click(force=True)
                page.wait_for_timeout(2000)
                page.get_by_text("Dispatch waybill query").first.click(force=True)
                page.wait_for_timeout(6000)
            except Exception:
                pass
            input_locator = encontrar_input(page)

        if input_locator is None:
            print(f"[{sigla}] INFO: Campo de AWB ainda não encontrado. Tentando em frames...")
            for frame in page.frames:
                try:
                    input_locator = encontrar_input(frame)
                    if input_locator is not None:
                        break
                except Exception:
                    continue

        if input_locator is None:
            try:
                if page.locator('.close-icon').count() and page.locator('.close-icon').first.is_visible():
                    page.locator('.close-icon').first.click(force=True)
                    page.wait_for_timeout(1000)
            except Exception:
                pass

        if input_locator is None:
            raise Exception(f"[{sigla}] FALHA: campo de AWB não encontrado na tela de destinatário.")

        input_locator.scroll_into_view_if_needed()
        input_locator.click(force=True)
        page.wait_for_timeout(500)

        awb_text = "\n".join(map(str, awb_lote))
        input_locator.fill(awb_text)
        page.keyboard.press("Enter")
        page.wait_for_timeout(5000)

        search_button = page.locator('button:has-text("Search"), button:has-text("Pesquisar"), button:has-text("Query")').first
        if search_button.count() and search_button.is_visible():
            search_button.click(force=True)
            page.wait_for_timeout(5000)

        print(f"[{sigla}] INFO: Consultando lote {indice_lote + 1} com {len(awb_lote)} AWBs...")
        try:
            page.locator('table').first.wait_for(state="visible", timeout=20000)
        except Exception:
            print(f"[{sigla}] WARN: Resultado em tabela não encontrado para o lote {indice_lote + 1}. Tentando continuar mesmo assim...")

        export_candidates = [
            '.ImileActionButton-root:has-text("Export")',
            '.ImileActionButton-root:has-text("Extrair")',
            'button:has-text("Export")',
            'button:has-text("Exportar")'
        ]
        export_clicked = False
        for selector in export_candidates:
            try:
                button = page.locator(selector).first
                if button.count() and button.is_visible():
                    button.click(force=True)
                    export_clicked = True
                    page.wait_for_timeout(1000)
                    break
            except Exception:
                continue

        if export_clicked:
            try:
                menu_item = page.locator('li.ImileMenuItem-root:has-text("Export All"), li.ImileMenuItem-root:has-text("Exportar tudo")').first
                if menu_item.count() and menu_item.is_visible():
                    menu_item.click(force=True)
                    page.wait_for_timeout(1000)
            except Exception:
                pass

            try:
                confirm_button = page.locator('.export-button').first
                if confirm_button.count() and confirm_button.is_visible():
                    confirm_button.click(force=True)
                    page.wait_for_timeout(1000)
                    print(f"[{sigla}] INFO: Abrindo a central de downloads para baixar o resultado do destinatário...")
                    page.locator('span.Imile-ButtonIcon-root svg path[d*="M8 0C3.57"]').locator('..').first.click(force=True)
                    page.wait_for_timeout(1000)
                    btn_baixar = page.locator('button:has-text("download")').last
                    btn_baixar.wait_for(state="visible", timeout=30000)
                    with page.expect_download(timeout=120000) as download_info:
                        btn_baixar.click(force=True)
                    download = download_info.value
                    caminho_destino = f"destinatario_{sigla}_lote_{indice_lote + 1}.xlsx"
                    download.save_as(caminho_destino)
                    print(f"[{sigla}] SUCCESS: Arquivo de destinatário salvo como '{caminho_destino}'")
                    return caminho_destino
            except Exception as e:
                print(f"[{sigla}] WARN: Falha ao exportar arquivo de destinatário no lote {indice_lote + 1}: {e}")

        print(f"[{sigla}] INFO: Export não disponível para o lote {indice_lote + 1}, tentando ler tabela diretamente da página...")
        df_html = carregar_tabela_html_para_df(page, sigla)
        if df_html is not None:
            temp_path = f"destinatario_{sigla}_lote_{indice_lote + 1}_html.xlsx"
            try:
                df_html.to_excel(temp_path, index=False)
                print(f"[{sigla}] SUCCESS: Tabela de destinatário gerada em '{temp_path}'")
                return temp_path
            except Exception as e:
                print(f"[{sigla}] WARN: Não foi possível salvar tabela HTML como Excel: {e}")
        return None

    arquivos_lotes = []
    for indice_lote, lote in enumerate(dividir_em_lotes(awb_list)):
        caminho_lote = consultar_lote(lote, indice_lote)
        if caminho_lote and os.path.exists(caminho_lote):
            arquivos_lotes.append(caminho_lote)

    if not arquivos_lotes:
        return None

    dfs_lotes = []
    for caminho_lote in arquivos_lotes:
        try:
            df_lote = pd.read_excel(caminho_lote)
            if df_lote is not None and not df_lote.empty:
                dfs_lotes.append(df_lote)
        except Exception as e:
            print(f"[{sigla}] WARN: Não foi possível ler o arquivo do lote '{caminho_lote}': {e}")

    if not dfs_lotes:
        return None

    df_consolidado = pd.concat(dfs_lotes, ignore_index=True)
    caminho_destino = f"destinatario_{sigla}_completo.xlsx"
    df_consolidado.drop_duplicates().to_excel(caminho_destino, index=False)
    print(f"[{sigla}] SUCCESS: Arquivo consolidado de destinatário salvo como '{caminho_destino}'")

    for caminho_lote in arquivos_lotes:
        try:
            if os.path.exists(caminho_lote):
                os.remove(caminho_lote)
        except Exception:
            pass

    return caminho_destino


def enriquecer_df_destinatario(df_base, df_dest, sigla):
    col_awb_base = buscar_coluna_por_tokens(df_base.columns, ['waybill', 'tracking', 'awb', 'bill'])
    col_awb_dest = buscar_coluna_por_tokens(df_dest.columns, ['waybill', 'tracking', 'awb', 'bill'])

    if not col_awb_base or not col_awb_dest:
        print(f"[{sigla}] WARN: Coluna AWB não localizada para enrichimento. Base: {col_awb_base}, Destino: {col_awb_dest}")
        return df_base

    colunas_desejadas = identificar_colunas_destinatario(df_dest)
    if not colunas_desejadas:
        print(f"[{sigla}] WARN: Nenhuma coluna de destinatário encontrada no resultado da pesquisa.")
        return df_base

    df_dest = df_dest[[col_awb_dest] + list(colunas_desejadas.values())].copy()
    df_dest.rename(columns={col_awb_dest: col_awb_base, **{v: k for k, v in colunas_desejadas.items()}}, inplace=True)
    df_dest[col_awb_base] = df_dest[col_awb_base].astype(str).str.strip()
    df_base[col_awb_base] = df_base[col_awb_base].astype(str).str.strip()

    df_consolidado = df_base.merge(df_dest.drop_duplicates(subset=[col_awb_base]), on=col_awb_base, how='left')
    print(f"[{sigla}] INFO: Enriquecimento de destinatário concluído. Total linhas base: {len(df_base)}")
    return df_consolidado


def enriquecer_com_destinatario(page, df_base, sigla):
    col_awb_base = buscar_coluna_por_tokens(df_base.columns, ['waybill', 'tracking', 'awb', 'bill'])
    if not col_awb_base:
        print(f"[{sigla}] WARN: Não foi possível encontrar coluna de AWB na base para enriquecer destinatário.")
        return df_base

    awbs = df_base[col_awb_base].dropna().astype(str).str.strip().unique().tolist()
    caminho_destinatario = buscar_arquivo_destinatario(page, awbs, sigla)
    if not caminho_destinatario or not os.path.exists(caminho_destinatario):
        print(f"[{sigla}] WARN: Não foi possível obter o arquivo de destinatário. Nenhum enriquecimento aplicado.")
        return df_base

    try:
        df_dest = pd.read_excel(caminho_destinatario)
    except Exception as e:
        print(f"[{sigla}] ERROR: Falha ao ler o arquivo de destinatário '{caminho_destinatario}': {e}")
        return df_base

    df_result = enriquecer_df_destinatario(df_base, df_dest, sigla)
    return df_result

# ==========================================
# FUNÇÃO DE DOWNLOAD
# ==========================================
def baixar_inventario_imile(page, sigla):
    """
    Navega no portal iMile e baixa o relatório de inventário para a base (sigla) logada.
    Retorna o caminho do arquivo temporário baixado.
    """
    print(f"[{sigla}] INFO: Navegando para a tela de 'Inventory Monitor'...")
    
    # Usamos try/except para cada clique para dar mais robustez e logs claros
    try:
        page.get_by_text("Monitor").first.click(force=True, timeout=20000)
    except Exception as e:
        print(f"[{sigla}] ERROR: Não foi possível clicar no menu 'Monitor'. Tentando alternativa... Detalhe: {e}")
        # Se o menu principal falha, às vezes um F5 e tentar de novo ajuda.
        page.reload(wait_until="domcontentloaded")
        page.get_by_text("Monitor").first.click(force=True, timeout=20000)
        
    page.wait_for_timeout(2000) # Pausa curta para o submenu abrir
    
    try:
        page.get_by_text("Operation Monitor").first.click(force=True)
    except Exception:
        # Tenta pelo link direto se o menu não funcionar
        print(f"[{sigla}] WARN: Clique em 'Operation Monitor' falhou, tentando acessar pelo URL.")
        current_url = page.url
        base_url = current_url.split('/#/') [0]
        page.goto(f"{base_url}/#/operation-monitor/inventory-monitor", wait_until="domcontentloaded")


    page.wait_for_timeout(2000)
    
    try:
        page.get_by_text("Inventory Monitor").first.click(force=True)
    except Exception:
        print(f"[{sigla}] WARN: Clique em 'Inventory Monitor' falhou, assumindo que já estamos na página correta.")
    
    # Espera a página de inventário carregar um elemento chave
    print(f"[{sigla}] INFO: Aguardando a página de inventário carregar...")
    page.locator('.ImileActionButton-root:has-text("Export")').first.wait_for(state="visible", timeout=45000)

    print(f"[{sigla}] INFO: Solicitando a exportação do inventário completo...")
    page.locator('.ImileActionButton-root:has-text("Export")').first.click(force=True)
    page.wait_for_timeout(1000)
    page.locator('li.ImileMenuItem-root:has-text("Export All")').first.click(force=True)
    page.wait_for_timeout(1000)
    
    # Clica no botão de confirmação dentro do modal de exportação, usando o seletor de classe mais robusto
    page.locator('.export-button').first.click(force=True)

    print(f"[{sigla}] INFO: Aguardando a iMile gerar o arquivo (isso pode demorar)...")
    page.wait_for_timeout(15000) # Espera estendida para a geração

    print(f"[{sigla}] INFO: Abrindo a central de downloads para baixar o arquivo...")
    # O ícone de download na barra de navegação superior
    page.locator('span.Imile-ButtonIcon-root svg path[d*="M8 0C3.57"]').locator('..').first.click(force=True)

    # Espera o último botão de download aparecer na lista
    btn_baixar = page.locator('button:has-text("download")').last
    btn_baixar.wait_for(state="visible", timeout=30000)

    # Inicia o download
    print(f"[{sigla}] INFO: Baixando o arquivo de inventário...")
    with page.expect_download(timeout=120000) as download_info:
        btn_baixar.click(force=True)

    download = download_info.value
    # Salva com um nome temporário para não confundir com arquivos finais
    caminho_temporario = f"temp_inventario_{sigla}.xlsx"
    download.save_as(caminho_temporario)

    print(f"[{sigla}] SUCCESS: Inventário baixado e salvo como '{caminho_temporario}'")
    return caminho_temporario

# ==========================================
# FUNÇÃO PARA PREPARAR O DATAFRAME DE UMA BASE
# ==========================================
def preparar_dataframe_base(caminho_origem, sigla_base):
    """
    Lê um arquivo de inventário, adiciona a coluna 'Base' e o retorna.
    Não faz nenhuma operação de escrita em arquivo.
    """
    if not os.path.exists(caminho_origem):
        print(f"ERROR: O arquivo de origem '{caminho_origem}' não foi encontrado.")
        return None

    try:
        print(f"INFO: Lendo dados do novo inventário de '{caminho_origem}'...")
        df_novo = pd.read_excel(caminho_origem)
        
        df_novo['Base'] = sigla_base
        
        colunas = ['Base'] + [col for col in df_novo.columns if col != 'Base']
        df_novo = df_novo[colunas]
        
        print(f"INFO: DataFrame para a base {sigla_base} preparado na memória.")
        return df_novo
        
    except Exception as e:
        print(f"ERROR: Falha ao processar o arquivo Excel de origem '{caminho_origem}': {e}")
        return None

# ==========================================
# FUNÇÃO FINAL PARA CONSOLIDAR, CALCULAR AGING E SALVAR
# ==========================================
def consolidar_calcular_e_salvar(lista_dfs, caminho_destino):
    """
    Recebe uma lista de DataFrames, consolida, calcula o aging e salva o resultado final.
    """
    if not lista_dfs:
        print("WARN: Nenhum dado foi coletado para processar. O arquivo final não será gerado.")
        return

    print("\nINFO: Etapa final: Consolidando todos os dados...")
    df_final = pd.concat(lista_dfs, ignore_index=True)
    df_final = adicionar_coluna_abrangencia(df_final)
    print(f"INFO: Consolidação concluída. Total de {len(df_final)} registros.")

    nome_aba_dados = "Sheet1"
    target_col_name = 'Last scan time'

    # --- Cálculo de Aging ---
    if target_col_name not in df_final.columns:
        print(f"WARN: A coluna '{target_col_name}' não foi encontrada. O 'aging' não será calculado.")
    else:
        print("INFO: Calculando 'aging'...")
        hoje = pd.to_datetime("today").normalize()
        
        # Converte a coluna alvo para datetime e normaliza para meia-noite (corrige o bug do -1)
        df_final[target_col_name] = pd.to_datetime(df_final[target_col_name], errors='coerce').dt.normalize()
        
        # Calcula a diferença em dias
        aging_values = (hoje - df_final[target_col_name]).dt.days
        
        # Renomeia a coluna e preenche com os valores calculados
        df_final.rename(columns={target_col_name: 'aging'}, inplace=True)
        df_final['aging'] = aging_values
        print("INFO: Coluna 'aging' calculada com sucesso.")

    # --- Salvamento Final ---
    try:
        print(f"INFO: Salvando resultado final em '{caminho_destino}', aba '{nome_aba_dados}'...")
        
        df_final = preparar_colunas_para_arquivo_excel(df_final)

        caminho_tmp = caminho_destino.replace('.xlsx', '_tmp.xlsx')

        # Salva como texto para evitar notação científica
        with pd.ExcelWriter(caminho_tmp, engine='openpyxl', mode='w') as writer:
            df_final.to_excel(writer, sheet_name=nome_aba_dados, index=False)
        
        # Formata waybill como texto
        formatar_waybill_como_texto(caminho_tmp)
        
        # Substitui atomicamente
        os.replace(caminho_tmp, caminho_destino)

        print(f"SUCCESS: Planilha '{caminho_destino}' gerada com sucesso (waybill em formato texto).")

    except PermissionError:
        print(f"FATAL: PERMISSÃO NEGADA para escrever no arquivo '{caminho_destino}'. Certifique-se de fechar o Excel.")
    except Exception as e:
        print(f"FATAL: Falha inesperada ao salvar a planilha final: {e}")


# ==========================================
# SCRIPT PRINCIPAL
# ==========================================
def main():
    print("======================================================")
    print("🚀 INICIANDO ATUALIZADOR DE LATÊNCIA IMILE 🚀")
    print("======================================================")
    
    lista_dataframes = [] # Lista para guardar os dataframes de cada base

    with sync_playwright() as p:
        for sigla, usuario, senha in BASES_SLA:
            caminho_temporario = None
            print(f"\n--- Processando base: {sigla} ---")
            
            if not usuario or not senha:
                print(f"WARN: Credenciais para '{sigla}' não encontradas no .env. Pulando...")
                continue

            browser = p.chromium.launch(headless=True, args=["--start-maximized"])
            context = browser.new_context(no_viewport=True, accept_downloads=True)
            page = context.new_page()

            try:
                login_imile(page, usuario, senha)
                caminho_temporario = baixar_inventario_imile(page, sigla)

                if caminho_temporario:
                    df_base = preparar_dataframe_base(caminho_temporario, sigla)
                    if df_base is not None:
                        df_base = enriquecer_com_destinatario(page, df_base, sigla)
                        lista_dataframes.append(df_base)
                else:
                    print(f"ERROR: Download para a base '{sigla}' falhou.")

            except Exception as e:
                print(f"FATAL: Ocorreu um erro crítico no processamento da base {sigla}: {e}")
                page.screenshot(path=f"error_{sigla}.png")
                print(f"INFO: Um print do erro foi salvo como 'error_{sigla}.png'")

            finally:
                print(f"[{sigla}] INFO: Finalizando e limpando...")
                context.close()
                browser.close()
                if caminho_temporario and os.path.exists(caminho_temporario):
                    os.remove(caminho_temporario)
                    print(f"[{sigla}] INFO: Arquivo temporário '{caminho_temporario}' removido.")
        
    # --- Etapa Final: Consolidar, Calcular Aging e Salvar ---
    if BASES_SLA:
        primeira_sigla = BASES_SLA[0][0]
        arquivo_consolidado_final = MAPA_ARQUIVOS_LATENCIA.get(primeira_sigla)
        if arquivo_consolidado_final:
             consolidar_calcular_e_salvar(lista_dataframes, arquivo_consolidado_final)

    print("\n======================================================")
    print("✅ Processo concluído para todas as bases.")
    print("======================================================")


if __name__ == "__main__":
    main()