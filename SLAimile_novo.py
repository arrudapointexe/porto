import os
import re
import pandas as pd
import unicodedata
from datetime import datetime, timedelta
from dotenv import load_dotenv
from playwright.sync_api import sync_playwright
from imile_utils import login_imile
from config import (BACKLOG_CRITICO, BACKLOG_DIA_DE_SEMANA, BACKLOG_FIM_DE_SEMANA,
                    ITR_LOC_CITIES, ITR_INT_CITIES, JML_LOC_CITIES, JML_INT_CITIES,
                    BASES_SLA)

load_dotenv()

# Mapear sigla para cidades LOC e INT
MAPA_CIDADES = {
    "ITR": {"LOC": ITR_LOC_CITIES, "INT": ITR_INT_CITIES},
    "JML": {"LOC": JML_LOC_CITIES, "INT": JML_INT_CITIES}
}

# ==============================================================
# FUNÇÕES AUXILIARES
# ==============================================================
def normalizar_cidade(texto):
    """Normaliza nome de cidade para comparação."""
    if pd.isna(texto):
        return ""
    texto = unicodedata.normalize('NFKD', str(texto))
    texto = ''.join(char for char in texto if not unicodedata.combining(char))
    texto = texto.encode('ascii', 'ignore').decode('ascii')
    texto = texto.strip().lower()
    texto = re.sub(r'[^a-z0-9\s]', ' ', texto)
    texto = re.sub(r'\s+', ' ', texto).strip()
    return texto

def buscar_coluna_por_tokens(colunas, tokens):
    """Encontra uma coluna pelo token."""
    tokens = [t.lower() for t in tokens]
    for col in colunas:
        col_lower = str(col).lower()
        if any(token in col_lower for token in tokens):
            return col
    return None

def eh_fim_de_semana():
    """Retorna True se hoje é sábado (5) ou domingo (6)."""
    return datetime.now().weekday() >= 5

def obter_backlog_limite():
    """Retorna o backlog máximo para considerar um pacote como 'vence hoje'."""
    if eh_fim_de_semana():
        return BACKLOG_FIM_DE_SEMANA
    else:
        return BACKLOG_DIA_DE_SEMANA

def classificar_abrangencia(cidade_nome, sigla):
    """Classifica uma cidade como LOC ou INT baseado no config."""
    cidade_norm = normalizar_cidade(cidade_nome)
    if not cidade_norm:
        return None
    
    cidades_map = MAPA_CIDADES.get(sigla, {})
    loc_cidades = [normalizar_cidade(c) for c in cidades_map.get("LOC", [])]
    int_cidades = [normalizar_cidade(c) for c in cidades_map.get("INT", [])]
    
    if cidade_norm in loc_cidades:
        return "LOC"
    elif cidade_norm in int_cidades:
        return "INT"
    return None

# ==============================================================
# PARTE 1: BAIXAR E CACHEAR INVENTÁRIO (MANHÃ)
# ==============================================================
def baixar_e_cachear_inventario(sigla, usuario, senha):
    """
    Baixa o inventário da base na iMile e salva em cache para consultas posteriores.
    Retorna o caminho do arquivo de inventário cacheado.
    """
    print(f"\n{'='*60}")
    print(f"📥 BAIXANDO E CACHEANDO INVENTÁRIO - {sigla}")
    print(f"{'='*60}")
    
    arquivo_cache = f"inventario_cache_{sigla}.xlsx"
    
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(accept_downloads=True)
        page = context.new_page()
        
        try:
            print(f"[{sigla}] Fazendo login...")
            login_imile(page, usuario, senha)
            
            print(f"[{sigla}] Navegando para Inventory Monitor...")
            page.wait_for_timeout(10000)
            page.get_by_text("Monitor").first.click(force=True)
            page.wait_for_timeout(2000)
            page.get_by_text("Operation Monitor").first.click(force=True)
            page.wait_for_timeout(2000)
            page.get_by_text("Inventory Monitor").first.click(force=True)
            page.wait_for_timeout(8000)
            
            print(f"[{sigla}] Solicitando exportação do inventário...")
            page.locator('.ImileActionButton-root:has-text("Export")').first.click(force=True)
            page.wait_for_timeout(1000)
            page.locator('li.ImileMenuItem-root:has-text("Export All")').first.click(force=True)
            page.wait_for_timeout(1000)
            page.locator('.export-button').first.click(force=True)
            
            print(f"[{sigla}] Aguardando geração do arquivo...")
            page.wait_for_timeout(15000)
            
            print(f"[{sigla}] Abrindo centro de downloads...")
            page.locator('span.Imile-ButtonIcon-root svg path[d*="M8 0C3.57"]').locator('..').first.click(force=True)
            page.wait_for_timeout(2000)
            
            btn_download = page.locator('button:has-text("download")').last
            btn_download.wait_for(state="visible", timeout=30000)
            
            print(f"[{sigla}] Baixando inventário...")
            with page.expect_download(timeout=120000) as download_info:
                btn_download.click(force=True)
            
            download = download_info.value
            download.save_as(arquivo_cache)
            
            print(f"✅ [{sigla}] Inventário cacheado: {arquivo_cache}")
            return arquivo_cache
            
        except Exception as e:
            print(f"❌ Erro ao baixar inventário de {sigla}: {e}")
            import traceback
            traceback.print_exc()
            return None
        finally:
            browser.close()

# ==============================================================
# PARTE 2: GERAR RELATÓRIO DE VENCIMENTOS POR CIDADE
# ==============================================================
def gerar_relatorio_vencimentos_por_cidade(sigla, arquivo_inventario):
    """
    Abre o inventário cacheado e gera um relatório de vencimentos.
    Retorna DataFrame com colunas: Cidade, RECEBIDO NA BASE, INVENTORY, EM ROTA, ENTREGUE, Total
    """
    print(f"\n{'='*60}")
    print(f"📊 GERANDO RELATÓRIO DE VENCIMENTOS - {sigla}")
    print(f"{'='*60}")
    
    if not arquivo_inventario or not os.path.exists(arquivo_inventario):
        print(f"❌ Arquivo de inventário não encontrado: {arquivo_inventario}")
        return None
    
    try:
        df = pd.read_excel(arquivo_inventario)
        print(f"✅ Inventário carregado: {len(df)} registros")
        print(f"   Colunas: {list(df.columns)}")
        
        # Identificar colunas principais
        col_city = buscar_coluna_por_tokens(df.columns, ['consignee city', 'city'])
        col_status = buscar_coluna_por_tokens(df.columns, ['last scan type', 'status', 'last scan'])
        col_backlog = buscar_coluna_por_tokens(df.columns, ['backlog time', 'backlog', 'backlog(station)'])
        
        if not all([col_city, col_status, col_backlog]):
            print(f"❌ Colunas obrigatórias não encontradas!")
            print(f"   Procurando por:")
            print(f"   - Cidade: {col_city}")
            print(f"   - Status: {col_status}")
            print(f"   - Backlog: {col_backlog}")
            return None
        
        print(f"✅ Colunas identificadas:")
        print(f"   - Cidade: {col_city}")
        print(f"   - Status: {col_status}")
        print(f"   - Backlog: {col_backlog}")
        
        # Preparar dados
        df_clean = df[[col_city, col_status, col_backlog]].copy()
        df_clean.columns = ['cidade', 'status', 'backlog']
        df_clean = df_clean.dropna(subset=['cidade', 'status', 'backlog'])
        
        # Normalizar cidades
        df_clean['cidade_norm'] = df_clean['cidade'].apply(normalizar_cidade)
        
        # Classificar por abrangência
        df_clean['abrangencia'] = df_clean.apply(
            lambda row: classificar_abrangencia(row['cidade_norm'], sigla), axis=1
        )
        
        # Filtrar apenas cidades da abrangência
        df_clean = df_clean[df_clean['abrangencia'].notna()].copy()
        
        if df_clean.empty:
            print(f"ℹ️ Nenhum pacote na abrangência de {sigla}")
            return pd.DataFrame()
        
        # Determinar backlog limite
        backlog_config = obter_backlog_limite()
        print(f"ℹ️ Backlog limite: LOC={backlog_config.get('loc', 2)}, INT={backlog_config.get('int', 3)}")
        
        # Marcar pacotes que "vencem hoje"
        df_clean['backlog_num'] = pd.to_numeric(df_clean['backlog'], errors='coerce').fillna(0)
        df_clean['vence_hoje'] = df_clean.apply(
            lambda row: row['backlog_num'] == backlog_config.get(row['abrangencia'].lower(), 2),
            axis=1
        )
        
        # Filtrar apenas vencimentos de hoje
        df_vencimentos = df_clean[df_clean['vence_hoje']].copy()
        
        if df_vencimentos.empty:
            print(f"ℹ️ Nenhum pacote vencendo hoje para {sigla}")
            return pd.DataFrame()
        
        print(f"✅ {len(df_vencimentos)} pacotes vencendo hoje")
        
        # Normalizar status para agrupação
        def normalizar_status(status_str):
            status_lower = str(status_str).lower().strip()
            if 'received' in status_lower or 'recebido' in status_lower:
                return 'RECEBIDO NA BASE'
            elif 'deliver' in status_lower or 'em rota' in status_lower:
                return 'EM ROTA'
            elif 'delivered' in status_lower or 'entregue' in status_lower:
                return 'ENTREGUE'
            else:
                return 'INVENTORY'
        
        df_vencimentos['status_port'] = df_vencimentos['status'].apply(normalizar_status)
        
        # Criar tabela dinâmica
        tabela = pd.crosstab(
            df_vencimentos['cidade'],
            df_vencimentos['status_port'],
            margins=True,
            margins_name='Total Geral'
        )
        
        # Reordenar colunas
        colunas_ordem = ['RECEBIDO NA BASE', 'INVENTORY', 'EM ROTA', 'ENTREGUE']
        colunas_presentes = [col for col in colunas_ordem if col in tabela.columns]
        
        if 'Total Geral' in tabela.columns:
            colunas_presentes.append('Total Geral')
        
        tabela = tabela[[col for col in colunas_presentes if col in tabela.columns]]
        
        # Preencher NaNs com 0
        tabela = tabela.fillna(0).astype(int)
        
        print(f"\n{'='*60}")
        print(f"RESUMO DE VENCIMENTOS - {sigla}")
        print(f"{'='*60}")
        print(tabela.to_string())
        
        return tabela
        
    except Exception as e:
        print(f"❌ Erro ao gerar relatório de vencimentos: {e}")
        import traceback
        traceback.print_exc()
        return None

# ==============================================================
# PARTE 3: SALVAR RELATÓRIO EM EXCEL E FORMATAR
# ==============================================================
def salvar_relatorio_excel(tabela, sigla):
    """Salva o relatório em Excel com formatação."""
    if tabela is None or tabela.empty:
        print(f"⚠️ Nenhum dado para salvar em {sigla}")
        return None
    
    arquivo_saida = f"Vencimentos_{sigla}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    try:
        with pd.ExcelWriter(arquivo_saida, engine='openpyxl') as writer:
            tabela.to_excel(writer, sheet_name='Vencimentos')
            
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            workbook = writer.book
            worksheet = writer.sheets['Vencimentos']
            
            # Header
            header_fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True, size=11)
            
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Total row
            last_row = len(tabela) + 2
            total_fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
            total_font = Font(color="FFFFFF", bold=True)
            
            for cell in worksheet[last_row]:
                cell.fill = total_fill
                cell.font = total_font
            
            # Ajustar largura e alinhar
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value or '')) > max_length:
                            max_length = len(str(cell.value or ''))
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    except:
                        pass
                worksheet.column_dimensions[column_letter].width = max(max_length + 2, 15)
        
        print(f"✅ Relatório salvo: {arquivo_saida}")
        return arquivo_saida
        
    except Exception as e:
        print(f"❌ Erro ao salvar relatório: {e}")
        return None

def formatar_mensagem_telegram(tabela, sigla):
    """Formata a tabela para envio no Telegram."""
    if tabela is None or tabela.empty:
        return f"✅ Nenhum pacote vencendo hoje em {sigla}."
    
    msg = f"📊 *VENCIMENTOS HOJE - {sigla}* 📊\n"
    msg += f"━━━━━━━━━━━━━━━━━━━━━━\n"
    msg += f"⏰ {datetime.now().strftime('%d/%m/%Y %H:%M')}\n\n"
    
    total_geral = 0
    for cidade, row in tabela.iterrows():
        if cidade == 'Total Geral':
            continue
        total = int(row.get('Total Geral', 0)) if 'Total Geral' in row.index else int(row.sum())
        if total > 0:
            msg += f"🏙️ *{cidade}*: {total} pacotes\n"
            total_geral += total
    
    msg += f"━━━━━━━━━━━━━━━━━━━━━━\n"
    msg += f"📦 *TOTAL:* {total_geral} pacotes\n"
    
    return msg

# ==============================================================
# ORQUESTRAÇÃO: ROTINA MATINAL (DOWNLOAD + CACHE)
# ==============================================================
def rotina_matinal_sla():
    """
    Executada uma vez pela manhã para atualizar o cache de inventários.
    Independente do schedule, o usuário pode rodar manualmente de manhã.
    """
    print(f"\n{'='*70}")
    print(f"🌅 ROTINA MATINAL - ATUALIZANDO CACHE DE INVENTÁRIOS")
    print(f"{'='*70}")
    
    for sigla, usuario, senha in BASES_SLA:
        if not usuario or not senha:
            print(f"⚠️ Credenciais ausentes para {sigla}")
            continue
        
        try:
            arquivo = baixar_e_cachear_inventario(sigla, usuario, senha)
            if arquivo:
                print(f"✅ Cache atualizado para {sigla}: {arquivo}\n")
        except Exception as e:
            print(f"❌ Erro ao processar {sigla}: {e}\n")

# ==============================================================
# ORQUESTRAÇÃO: ROTINA DE CONSULTA (12h, 15h, 18h, 21h)
# ==============================================================
def rotina_consulta_sla():
    """
    Executada 4x ao dia (12h, 15h, 18h, 21h) para gerar relatórios
    de vencimentos consultando o cache de inventários.
    """
    print(f"\n{'='*70}")
    print(f"⏰ ROTINA DE CONSULTA SLA - {datetime.now().strftime('%H:%M')}")
    print(f"{'='*70}")
    
    for sigla, usuario, senha in BASES_SLA:
        print(f"\n▶️ Processando {sigla}...")
        
        arquivo_cache = f"inventario_cache_{sigla}.xlsx"
        
        if not os.path.exists(arquivo_cache):
            print(f"⚠️ Cache não encontrado para {sigla}. Pulando...")
            print(f"   Dica: Execute rotina_matinal_sla() de manhã para atualizar o cache.")
            continue
        
        try:
            # Gerar relatório
            tabela = gerar_relatorio_vencimentos_por_cidade(sigla, arquivo_cache)
            
            if tabela is not None and not tabela.empty:
                # Salvar em Excel
                arquivo_saida = salvar_relatorio_excel(tabela, sigla)
                
                # Formatar mensagem
                mensagem = formatar_mensagem_telegram(tabela, sigla)
                
                print(f"\n{mensagem}")
                
                # Aqui o botTelegram.py vai chamar essa função e enviar a mensagem
                # Retornando para integração com o bot
                return {
                    "sigla": sigla,
                    "tabela": tabela,
                    "arquivo": arquivo_saida,
                    "mensagem": mensagem
                }
            else:
                print(f"ℹ️ Nenhum vencimento para {sigla} neste momento.")
        
        except Exception as e:
            print(f"❌ Erro ao processar {sigla}: {e}")
            import traceback
            traceback.print_exc()
    
    return None

# ==============================================================
# FUNÇÃO WRAPPER PARA USO DO BOT
# ==============================================================
def executar_rotina_sla_do_bot(tipo="consulta"):
    """
    Função wrapper chamada pelo botTelegram.py.
    tipo: "matinal" ou "consulta"
    """
    if tipo == "matinal":
        rotina_matinal_sla()
        return "✅ Cache de inventários atualizado com sucesso!"
    elif tipo == "consulta":
        resultado = rotina_consulta_sla()
        if resultado:
            return resultado
        else:
            return None
    return None

if __name__ == "__main__":
    # Para testes
    rotina_matinal_sla()
