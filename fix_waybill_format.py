#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Script para corrigir formato de waybills em arquivos Excel.
Converte notação científica para texto.
"""

from openpyxl import load_workbook
import os

def formatar_waybill_como_texto(caminho_arquivo):
    """
    Formata todas as colunas de waybill/tracking como texto no Excel para evitar notação científica.
    """
    if not os.path.exists(caminho_arquivo):
        print(f"Erro: Arquivo não encontrado: {caminho_arquivo}")
        return False
    
    try:
        print(f"Abrindo arquivo: {caminho_arquivo}")
        wb = load_workbook(caminho_arquivo)
        ws = wb.active
        
        print(f"Folha ativa: {ws.title}")
        print(f"Colunas: {[cell.value for cell in ws[1]]}")
        
        colunas_formatadas = []
        
        # Itera pelas colunas do header para identificar waybill/tracking
        for col_idx, col_cell in enumerate(ws[1], 1):
            if col_cell.value and any(x in str(col_cell.value).lower() for x in ['waybill', 'tracking', 'order', 'pack']):
                col_name = col_cell.value
                print(f"\nFormatando coluna {col_idx} ({col_name})...")
                colunas_formatadas.append(col_name)
                
                # Formata todas as células dessa coluna como texto
                count = 0
                for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx), start=2):
                    for cell in row:
                        cell.number_format = '@'  # @ = texto no Excel
                        if cell.value is not None:
                            # Converte para string e remove espaços
                            cell.value = str(cell.value).strip()
                            count += 1
                
                print(f"  {count} células formatadas como texto")
        
        if not colunas_formatadas:
            print("Nenhuma coluna de waybill/tracking encontrada!")
            return False
        
        print(f"\nSalvando arquivo...")
        wb.save(caminho_arquivo)
        print(f"✓ Arquivo '{os.path.basename(caminho_arquivo)}' formatado com sucesso!")
        print(f"  Colunas formatadas: {', '.join(colunas_formatadas)}")
        return True
        
    except Exception as e:
        print(f"Erro ao formatar arquivo: {e}")
        import traceback
        traceback.print_exc()
        return False

if __name__ == "__main__":
    os.chdir(r"c:\bots")
    
    # Formata os três arquivos
    arquivos = [
        "latencia_consolidada.xlsx",
        "latenciaJML.xlsx",
        "latenciaITR.xlsx"
    ]
    
    print("=" * 60)
    print("FORMATADOR DE WAYBILL - Convertendo notação científica")
    print("=" * 60)
    
    for arquivo in arquivos:
        if os.path.exists(arquivo):
            print(f"\n{'=' * 60}")
            formatar_waybill_como_texto(arquivo)
        else:
            print(f"Arquivo não encontrado: {arquivo}")
    
    print(f"\n{'=' * 60}")
    print("✓ Processo concluído!")
